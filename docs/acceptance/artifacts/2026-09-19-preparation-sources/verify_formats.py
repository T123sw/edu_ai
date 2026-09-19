"""Opt-in synthetic DOCX/scanned PDF/PPTX MinerU verification. Requires python-docx, python-pptx, Pillow. Credentials are read only from --env-file and never written."""
from pathlib import Path
import json, time, io
import dotenv
import argparse
arguments=argparse.ArgumentParser()
arguments.add_argument('--env-file',required=True)
arguments.add_argument('--output-dir',required=True)
args=arguments.parse_args()
settings=dotenv.dotenv_values(args.env_file)
dotenv.load_dotenv=lambda *a,**kw:False
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont
import httpx
from app.integrations.pdf.mineru_cloud import MinerUCloudProvider
from app.course_preparation.adapters.documents import DocumentParser
out=Path(args.output_dir); out.mkdir(parents=True,exist_ok=False)
font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',30)
phrases=['TOOLS NEED VALIDATED ARGUMENTS','FAILURE RETURNS INVALID_ARGUMENT','SUCCESS ADDS 12 AND 30 TO GET 42','FINAL PAGE REQUIRES API AUTHENTICATION']
image=Image.new('RGB',(1000,600),'white');draw=ImageDraw.Draw(image)
draw.text((40,40),'TOOL FLOW DIAGRAM',fill='black',font=font)
for i,label in enumerate(['MODEL','VALIDATE','TOOL']):
 x=40+300*i;draw.rectangle((x,180,x+250,300),outline='black',width=3);draw.text((x+15,220),label,fill='black',font=font)
image.save(out/'diagram.png')
doc=Document();doc.add_heading('Calculator Agent Reference',0)
for phrase in phrases:doc.add_paragraph(phrase)
table=doc.add_table(rows=2,cols=2);table.cell(0,0).text='Input';table.cell(0,1).text='Expected'
table.cell(1,0).text='12 + 30';table.cell(1,1).text='42'
doc.add_picture(str(out/'diagram.png'),width=__import__('docx').shared.Inches(5));doc.save(out/'reference.docx')
slides=Presentation()
for i in range(2):
 slide=slides.slides.add_slide(slides.slide_layouts[6]);box=slide.shapes.add_textbox(Inches(.3),Inches(.3),Inches(9),Inches(2))
 box.text='Calculator Agent Slide '+str(i+1)+'\n'+'\n'.join(phrases[i*2:i*2+2])
 if i==0:slide.shapes.add_picture(str(out/'diagram.png'),Inches(.3),Inches(2.4),width=Inches(8))
 else:
  table=slide.shapes.add_table(2,2,Inches(.5),Inches(3),Inches(6),Inches(1.5)).table
  table.cell(0,0).text='INPUT';table.cell(0,1).text='RESULT';table.cell(1,0).text='12 + 30';table.cell(1,1).text='42'
slides.save(out/'reference.pptx')
pages=[]
for i in range(2):
 page=Image.new('RGB',(1600,2000),'white');d=ImageDraw.Draw(page)
 for j,phrase in enumerate(phrases[i*2:i*2+2]):d.text((70,100+j*150),phrase,fill='black',font=font)
 d.text((70,500),'Scanned source page '+str(i+1),fill='black',font=font)
 if i==0:page.paste(image,(70,700))
 pages.append(page)
pages[0].save(out/'reference.pdf',save_all=True,append_images=pages[1:])
report=[]
for suffix in ('docx','pdf','pptx'):
 calls=[];tasks=[]
 class Client(httpx.Client):
  def request(self,method,url,**kw):
   started=time.monotonic()
   try:
    response=super().request(method,url,**kw)
    calls.append({'method':method,'phase':'batch' if 'file-urls' in str(url) else 'poll' if 'extract-results' in str(url) else 'upload' if method=='PUT' else 'download','status':response.status_code,'seconds':round(time.monotonic()-started,3)})
    return response
   except Exception as exc:
    calls.append({'method':method,'error':type(exc).__name__,'seconds':round(time.monotonic()-started,3)})
    raise
 provider=MinerUCloudProvider(api_key=settings['PDF_MINERU_CLOUD_API_KEY'],base_url=settings.get('PDF_MINERU_CLOUD_BASE_URL'),client_factory=Client)
 started=time.monotonic();result={'format':suffix}
 try:
  parsed=DocumentParser(mineru_factory=lambda:provider).parse((out/f'reference.{suffix}').read_bytes(),filename=f'reference.{suffix}',deadline=time.monotonic()+240,on_remote_task=tasks.append)
  (out/f'{suffix}-extracted.md').write_text(parsed.content)
  normalized=' '.join(parsed.content.upper().split())
  result.update(status='parsed',characters=len(parsed.content),source_hash=parsed.source_hash,defects=parsed.defects,assets=len(parsed.assets),units=len(parsed.units),checks={phrase:phrase in normalized for phrase in phrases})
 except Exception as exc:result.update(status='failed',error_code=getattr(exc,'code',type(exc).__name__))
 result.update(seconds=round(time.monotonic()-started,3),calls=calls,remote_tasks=tasks)
 report.append(result);(out/'results.json').write_text(json.dumps(report,indent=2))
 print(json.dumps(result),flush=True)
